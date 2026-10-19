#!/usr/bin/env python3
"""tests/quick/test_book_formulas.py — каждая формула каждой декларации книги считает измеримый результат.

Книги не перечисляются: каталог — резолвер сервера (`server_tables`), декларации — дверь книг; своя книга проекта и новая
формула попадают в замер сами. На книгу — один пересчёт: восемь строк данных, четыре рода входов (текст, число строкой,
время «MM:SS», малые числа) и сдвиг по столбцу; формулу строки собирает узел (`core.formulas.compiled`), как у
материализатора; истина — пересчёт LibreOffice. Измерима формула без ошибки, давшая хотя бы два разных значения, не
запасных; снимок той же сущности (очередь, движок снимка) = книга клетка в клетку. Дыра формулы: каждая формула без
IFERROR на краевых входах своего типа и формата — ошибка на заполненных входах только из объявленных `degrade_covers`.
Запись снимка судит тип и формат: новая строка, правка ячейки и правка строки. Дверь — пробы с изъяном каждого рода:
отказ с путём столбца, законная проба — книга.
"""
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from core.contracts import ContractError
from core.declaration import declarations_root
from core.declaration_schemas import book_sheets
from core.engine import TableMaterializer, TableMaterializerError
from core.engine.table_materializer import book_defaults, book_schema
from core.engine.template_resolver import server_tables
from core.excel import ExcelEngine
from core.formulas import DATA_ROW, compiled, computed, declared_language
from core.ids import IDGenerator
from core.state import StateManager
from core.tables import TableEngine

_checks = 0
_fails: list[str] = []
ERRORS = ("#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
FAMILIES = ("текст", "число", "время", "малые")
ROWS = [(family, moved) for family in FAMILIES for moved in (False, True)]
DEFAULTS = book_defaults(server_tables(declarations_root()))[1]   # общие правила книг: форматы строк, что покрывает запасное


def check(label: str, cond: bool, detail: str = "") -> None:
    global _checks
    _checks += 1
    print(f"  {'✓' if cond else '✗'} {label}  {detail if not cond else ''}".rstrip())
    if not cond:
        _fails.append(label)


def value(col, i: int, family: str, moved: bool):
    """Вход столбца по типу и роду; `moved` — сдвиг, разный у каждого столбца: меняется и отношение двух столбцов. Строка
    с форматом — его образцы (иное запись снимка отбивает)."""
    if col.format:
        pool = DEFAULTS.formats[col.format].examples
        return pool[(i + 3 * moved + FAMILIES.index(family)) % len(pool)]
    kind = col.type
    if kind in ("integer", "float") and family == "малые":
        return (i % 3) + int(moved) if kind == "integer" else (i % 3) * 0.5 + 0.5 * moved
    if kind == "integer":
        return (7 + 4 * i) * 3 + i + 1 if moved else 7 + 4 * i
    if kind == "float":
        return (2.5 + i) * 3 + i + 1 if moved else 2.5 + i
    if kind == "enum":
        return list(col.enum)[1 if moved and len(col.enum) > 1 else 0] if col.enum else f"v{i}{int(moved)}"
    if kind == "boolean":
        return not moved
    if kind == "date":
        return f"2026-0{1 + i % 9}-1{int(moved)}"
    if kind == "list":
        return "a,b,c" if moved else "a,b"
    if family == "число":
        return str((7 + 4 * i) * 3 + i + 1 if moved else 7 + 4 * i)
    if family == "время":
        return f"0{2 if moved else 1}:{(13 + 7 * i) % 60:02d}"
    return " ".join(["слово"] * ((i % 4) + 1 + (3 if moved else 0)))


def measured(name: str, tables: Path, work: Path) -> dict[tuple[str, str], tuple[list, list, object]]:
    """Значения формул по строкам родов: книги после пересчёта LibreOffice и снимка той же сущности после применения
    очереди (движок снимка), и запасное значение столбца. Строки снимка — настоящим путём: `append` в очередь и
    `execute_queue`."""
    (work / name).mkdir()
    (work / name / "read.json").write_text("{}", encoding="utf-8")
    engine, snapshots = ExcelEngine(work), TableEngine(StateManager(work), IDGenerator())
    materializer = TableMaterializer(engine, tables, snapshots)
    book = work / name / f"{name}.xlsx"
    materializer.materialize(name, f"{name}/{name}.xlsx")
    schema = book_schema(tables, name)
    _, language = declared_language(TableMaterializerError)
    sheets, rules = book_sheets(schema), materializer._degrade_rules(schema)
    wb = openpyxl.load_workbook(book)
    for sheet in schema.sheets:
        ws = wb[sheet.name]
        for at, (family, moved) in enumerate(ROWS, start=DATA_ROW):
            inputs = {}
            for i, col in enumerate(sheet.columns, start=1):
                if col.flag == "F" and col.formula:
                    formula = compiled(col.formula, sheet.name, sheets, at, language)
                    ws.cell(row=at, column=i, value=materializer._guarded(formula, col, rules))
                else:
                    given = value(col, i - 1, family, moved)
                    # Ключ строки (flag: id) лист держит одним на строку: номер строки делает его уникальным.
                    inputs[col.name] = f"{given}#{at}" if col.flag == "id" and given not in (None, "") else given
                    ws.cell(row=at, column=i, value=inputs[col.name])
            snapshots.append(name, sheet.name, inputs)
    wb.save(book)
    snapshots.execute_queue(name)
    snap = snapshots.load_snapshot(name)
    got = openpyxl.load_workbook(engine._recalc_via_lo(book), data_only=True)
    return {(sheet.name, col.name): ([got[sheet.name].cell(row=at, column=i).value
                                      for at in range(DATA_ROW, DATA_ROW + len(ROWS))],
                                     [row.get(col.name) for row in snap[sheet.name]["rows"].values()],
                                     materializer.fallback(col, rules))
            for sheet in schema.sheets for i, col in enumerate(sheet.columns, start=1)
            if col.flag == "F" and col.formula}


def same(book: object, snap: object) -> bool:
    """Значение книги и снимка одно: числа — с точностью пересчёта, пусто книги (None) — пустая строка снимка."""
    if isinstance(book, (int, float)) and isinstance(snap, (int, float)) and not isinstance(book, bool):
        return abs(book - snap) <= 1e-9 * max(1.0, abs(book))
    return ("" if book is None else str(book)) == ("" if snap is None else str(snap))


print("== 1. Каждая объявленная формула каждой книги — измеримый результат (истина — LibreOffice) ==")
TABLES = server_tables(declarations_root())
_work = Path(tempfile.mkdtemp(prefix="book_formulas_"))
try:
    for _path in sorted(TABLES.glob("*.schema.yaml")):
        _name = _path.name.removesuffix(".schema.yaml")
        try:
            _values = measured(_name, TABLES, _work)
        except ContractError as _e:   # нет движка, дверь или очередь отказали — красное, а не пропуск
            check(f"{_name}: книга и снимок собраны и посчитаны", False, f"{_e.code} {_e.facts} {_e.errors}")
            continue
        except ValueError as _e:   # `compiled` не собрал формулу, которую дверь пустила, — изъян узла
            check(f"{_name}: узел собрал каждую формулу, пущенную дверью", False, str(_e)[:200])
            continue
        for (_sheet, _col), (_got, _snap, _fallback) in _values.items():
            _errors = [v for v in _got if str(v).startswith(ERRORS)]
            _real = {str(v) for v in _got if v not in (_fallback, None, "") and not str(v).startswith(ERRORS)}
            check(f"{_name}.{_sheet}.{_col}: без ошибки, значений {len(_real)} (нужно ≥ 2)",
                  not _errors and len(_real) >= 2, f"ошибки {_errors[:2]}, значения {sorted(_real)[:3]}")
            _apart = [(b, s) for b, s in zip(_got, _snap) if not same(b, s)]
            check(f"{_name}.{_sheet}.{_col}: снимок = книга по {len(_got)} строкам",
                  len(_snap) == len(_got) and not _apart, f"расходятся {_apart[:3]}, строк снимка {len(_snap)}")
finally:
    shutil.rmtree(_work, ignore_errors=True)

print("== 1б. Строка снимка остаётся текстом: похожее на число, формулу, логическое и время движок не разбирает ==")
_TEXTS = ["007", "=1+1", "TRUE", "1e3", "01:13", " 5 ", "50%"]
_dir = Path(tempfile.mkdtemp(prefix="book_text_"))
try:
    (_dir / "t.schema.yaml").write_text(yaml.safe_dump({
        "book": "t", "level": "video", "source": "test",
        "sheets": [{"name": "S", "columns": [{"name": "s", "type": "string", "flag": "W"},
                                             {"name": "n", "type": "integer", "flag": "F", "formula": "=LEN(s)"}]}]},
        allow_unicode=True), encoding="utf-8")
    (_dir / "t").mkdir()
    (_dir / "t" / "read.json").write_text("{}", encoding="utf-8")
    _tables = TableEngine(StateManager(_dir), IDGenerator())
    TableMaterializer(ExcelEngine(_dir), _dir, _tables).materialize("t", "t/t.xlsx")
    for _text in _TEXTS:
        _tables.append("t", "S", {"s": _text})
    _tables.execute_queue("t")
    _lens = [row.get("n") for row in _tables.load_snapshot("t")["S"]["rows"].values()]
    for _text, _len in zip(_TEXTS, _lens):   # истина — длина строки Python, не движок
        check(f"LEN({_text!r}) = {len(_text)}", _len == len(_text), f"снимок дал {_len!r}")
    check(f"строк снимка {len(_lens)} = входов {len(_TEXTS)}", len(_lens) == len(_TEXTS))
except ContractError as _e:
    check("снимок с текстовыми входами собран и посчитан", False, f"{_e.code} {_e.facts}")
finally:
    shutil.rmtree(_dir, ignore_errors=True)

print("== 3. Формула без IFERROR на краевых входах своего типа: ошибка на заполненных входах — только покрытая ==")
# Края по объявленному типу; строке с форматом — его образцы. Пустой вход — неполные данные (запасное); иначе ошибка
# движка обязана быть из `degrade_covers` столбца или `_defaults.yaml`, другая — дыра формулы.
EDGES = {"integer": [0, -1, 1_000_000_000], "float": [0, -1.5, 1e9], "boolean": [False, True],
         "string": ["слово", "два  слова", "7", " "], "date": ["2026-12-31"], "list": ["a", "a,b,c"]}
_defaults = DEFAULTS
_, _language = declared_language(TableMaterializerError)


def edges(col) -> list:
    if col.format:
        return list(_defaults.formats[col.format].examples)
    if col.type == "enum":
        return list(col.enum) or ["v"]
    return EDGES[col.type]


def normal(col, i: int) -> object:
    """Обычный вход — свой у каждого столбца (два времени не равны, делитель не ноль)."""
    pool = edges(col)
    if col.type in ("integer", "float") and not col.format:
        return 10 + i
    return pool[(1 + i) % len(pool)]


for _path in sorted(TABLES.glob("*.schema.yaml")):
    _name = _path.name.removesuffix(".schema.yaml")
    _schema = book_schema(TABLES, _name)
    _inputs = [(s.name, c) for s in _schema.sheets for c in s.columns if not (c.flag == "F" and c.formula)]
    _plan = [None] + [(s, c.name, v) for s, c in _inputs for v in [None, *edges(c)]]
    _sheets = {}
    for _sheet in _schema.sheets:
        _rows = {}
        for _k, _change in enumerate(_plan):
            _row = {c.name: normal(c, i) for i, c in enumerate(_sheet.columns) if not (c.flag == "F" and c.formula)}
            if _change and _change[0] == _sheet.name:
                _row[_change[1]] = _change[2]
            _rows[f"R{_k}"] = {k: v for k, v in _row.items() if v is not None}
        _sheets[_sheet.name] = {"rows": _rows, **({"record": True} if _sheet.record else {}), "schema": {
            c.name: ({"type": c.type, "computed": True, "formula": c.formula, "fallback": None}
                     if c.flag == "F" and c.formula else {"type": c.type}) for c in _sheet.columns}}
    _raw = computed(_sheets, _language)
    for _sheet in _schema.sheets:
        for _col in [c for c in _sheet.columns if c.flag == "F" and c.formula]:
            _covers = set(_col.degrade_covers if _col.degrade_covers is not None else _defaults.degrade_covers)
            _holes = [(f"{c[0]}.{c[1]}={c[2]!r}" if c else "обычные входы", got)
                      for k, c in enumerate(_plan)
                      if isinstance(got := _raw[_sheet.name][f"R{k}"][_col.name], str) and got.startswith("#")
                      and not (c and c[2] is None) and got not in _covers]
            check(f"{_name}.{_sheet.name}.{_col.name}: {len(_plan)} строк, непокрытых ошибок {len(_holes)}",
                  not _holes, f"покрыто {sorted(_covers)}; дыры {_holes[:3]}")

print("== 3б. Запись снимка судит тип и формат столбца: чужое — VALIDATION_ERROR с полем и правилом, своё — в очередь ==")
_dir = Path(tempfile.mkdtemp(prefix="book_write_"))
try:
    (_dir / "_defaults.yaml").write_text((TABLES / "_defaults.yaml").read_text(encoding="utf-8"), encoding="utf-8")
    _fmt = sorted(_defaults.formats)[0]
    (_dir / "w.schema.yaml").write_text(yaml.safe_dump({
        "book": "w", "level": "video", "source": "test",
        "sheets": [{"name": "S", "columns": [
            {"name": "n", "type": "integer", "flag": "W"}, {"name": "x", "type": "float", "flag": "W"},
            {"name": "b", "type": "boolean", "flag": "W"}, {"name": "d", "type": "date", "flag": "W"},
            {"name": "t", "type": "string", "flag": "W", "format": _fmt}]}]}, allow_unicode=True), encoding="utf-8")
    (_dir / "w").mkdir()
    (_dir / "w" / "read.json").write_text("{}", encoding="utf-8")
    _tables = TableEngine(StateManager(_dir), IDGenerator())
    TableMaterializer(ExcelEngine(_dir), _dir, _tables).materialize("w", "w/w.xlsx")
    _example = _defaults.formats[_fmt].examples[0]
    for _case, _data, _rule in [("число словами в integer", {"n": "сто"}, "type"),
                                ("bool в integer", {"n": True}, "type"), ("текст в float", {"x": "1.5"}, "type"),
                                ("число в boolean", {"b": 1}, "type"), ("дата не ISO", {"d": "31.12.2026"}, "type"),
                                (f"строка вне формата {_fmt}", {"t": "слово"}, "format"),
                                (f"число вместо строки формата {_fmt}", {"t": 75}, "format"),
                                ("свои значения", {"n": 3, "x": 1.5, "b": False, "d": "2026-12-31", "t": _example}, None),
                                ("пусто — неполные данные", {"n": None, "t": ""}, None)]:
        try:
            _tables.append("w", "S", _data)
            _said = ("принято", "")
        except ContractError as _e:
            _said = (_e.code, str(_e.errors))
        if _rule is None:
            check(f"{_case}: в очередь", _said[0] == "принято", str(_said))
        else:
            check(f"{_case}: VALIDATION_ERROR, правило {_rule}",
                  _said[0] == "VALIDATION_ERROR" and f"'rule': '{_rule}'" in _said[1], str(_said))
    _tables.execute_queue("w")
    _id = next(iter(_tables.load_snapshot("w")["S"]["rows"]))
    for _case, _data, _rule in [("правка ячейки вне формата", ("t", "1:02:03"), "format"),
                                ("правка ячейки словами в integer", ("n", "три"), "type")]:
        try:
            _tables.set("w", "S", _id, _data[0], _data[1])
            _said = ("принято", "")
        except ContractError as _e:
            _said = (_e.code, str(_e.errors))
        check(f"{_case}: VALIDATION_ERROR, правило {_rule}",
              _said[0] == "VALIDATION_ERROR" and f"'rule': '{_rule}'" in _said[1], str(_said))
    try:
        _tables.update("w", "S", _id, {"t": "слово"})
        _said = ("принято", "")
    except ContractError as _e:
        _said = (_e.code, str(_e.errors))
    check("правка строки вне формата: VALIDATION_ERROR, правило format",
          _said[0] == "VALIDATION_ERROR" and "'rule': 'format'" in _said[1], str(_said))
except ContractError as _e:
    check("книга записи собрана", False, f"{_e.code} {_e.facts}")
finally:
    shutil.rmtree(_dir, ignore_errors=True)

print("== 2. Дверь книги: формула с изъяном — отказ с путём столбца, законная — книга ==")
_PROBES = {
    "законная": ("=a/b*100", None),
    "имя не столбец": ("=a/views", "formula_name"),
    "лист не книги": ("=a/NOPE.b", "formula_sheet"),
    "функция не объявлена": ("=FOO(a)", "formula_function"),
    "лист-список вне агрегата": ("=a/L.b", "formula_row"),
    "оператор без левого операнда": ("=IF(>=5,1,0)", "formula_syntax"),
    "скобка не закрыта": ("=(a+b", "formula_syntax"),
}
for _case, (_formula, _rule) in _PROBES.items():
    _dir = Path(tempfile.mkdtemp(prefix="book_door_"))
    try:
        (_dir / "p.schema.yaml").write_text(yaml.safe_dump({
            "book": "p", "level": "video", "source": "test",
            "sheets": [{"name": "S", "columns": [
                {"name": "a", "type": "integer", "flag": "W"}, {"name": "b", "type": "integer", "flag": "W"},
                {"name": "r", "type": "float", "flag": "F", "formula": _formula}]},
                {"name": "L", "columns": [{"name": "b", "type": "integer", "flag": "W"}]}]},
            allow_unicode=True), encoding="utf-8")
        try:
            book_schema(_dir, "p")
            _said = ("", "")
        except TableMaterializerError as _e:
            _said = (_e.code, str(_e.facts.get("details", "")))
        if _rule is None:
            check(f"{_case}: дверь пускает книгу", _said == ("", ""), str(_said))
        else:
            check(f"{_case}: отказ {_rule} с путём sheets.0.columns.2.formula",
                  _said[0] == "DECLARATION_INVALID" and "sheets.0.columns.2.formula" in _said[1]
                  and _rule in _said[1], str(_said))
    finally:
        shutil.rmtree(_dir, ignore_errors=True)

print(f"\n{'=' * 50}")
print(f"РЕЗУЛЬТАТ: {_checks - len(_fails)}/{_checks} прошло")
if _fails:
    print("ПРОВАЛЫ:")
    for _f in _fails:
        print(f"  - {_f}")
    sys.exit(1)
print("ВСЁ ЗЕЛЁНОЕ ✅")
